#!/usr/bin/env python3
"""
Ghi kết quả B1/B2 vào file Word báo cáo chung.

Thường KHÔNG cần chạy tay — run_b1_eval.py / run_b2_eval.py gọi tự động
khi GHI_BAO_CAO = True.

Chạy tay (nếu muốn ghi lại từ log cũ):
    conda activate phenikaa
    python scripts/ghi_bao_cao.py b1
    python scripts/ghi_bao_cao.py b2
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from _maptr_env import parse_metrics  # noqa: E402

# =============================================================================
# CONFIG — CHỈ SỬA PHẦN NÀY
# =============================================================================

PROJECT_ROOT = Path("/home/khanh247/Documents/Survey/phenikaa")

# File báo cáo chung (cả B0, B1, B2, B3…)
BAO_CAO_DOCX = PROJECT_ROOT / "docs/BAO_CAO_QUA_TRINH.docx"

# Bản copy cùng nội dung (giữ tên cũ cho tiện)
BAO_CAO_COPY = PROJECT_ROOT / "docs/BAO_CAO_QUA_TRINH_B0_B1.docx"

LOG_B1 = PROJECT_ROOT / "outputs/nusc_eval/e1_camera/eval_log.txt"
LOG_B2 = PROJECT_ROOT / "outputs/nusc_eval/e2_fusion/eval_log.txt"

# =============================================================================


def _fmt(v: float) -> str:
    return f"{v:.3f}"


def _write_metrics_md(output_dir: Path, title: str, metrics: dict[str, float], log_path: Path) -> None:
    md = output_dir / "metrics.md"
    md.write_text(
        f"# {title}\n\n"
        f"| Class | AP |\n|-------|-----|\n"
        f"| divider | {_fmt(metrics['divider'])} |\n"
        f"| ped_crossing | {_fmt(metrics['ped_crossing'])} |\n"
        f"| boundary | {_fmt(metrics['boundary'])} |\n"
        f"| **mAP** | **{_fmt(metrics['map'])}** |\n\n"
        f"Log: `{log_path.name}`\n",
        encoding="utf-8",
    )


def _update_docx(docx_path: Path, *, step: str, metrics: dict[str, float], when: str) -> None:
    from docx import Document

    doc = Document(str(docx_path))

    # Bảng tiến độ (table 1)
    if len(doc.tables) > 1:
        t = doc.tables[1]
        row_idx = 2 if step == "b1" else 3
        if row_idx < len(t.rows):
            t.rows[row_idx].cells[4].text = f"ĐÃ CHẠY {when}"

    if step == "b1" and len(doc.tables) > 2:
        t = doc.tables[2]
        vals = [metrics["divider"], metrics["ped_crossing"], metrics["boundary"], metrics["map"]]
        for i, val in enumerate(vals, start=1):
            t.rows[i].cells[1].text = _fmt(val)
        for para in doc.paragraphs:
            if para.text.startswith("Ngày chạy B1:"):
                para.text = f"Ngày chạy B1 (lần gần nhất): {when}"

    if step == "b2" and len(doc.tables) > 3:
        t = doc.tables[3]
        b2_vals = [metrics["divider"], metrics["ped_crossing"], metrics["boundary"], metrics["map"]]
        for i, val in enumerate(b2_vals, start=1):
            t.rows[i].cells[1].text = _fmt(val)

        # Điền cột AP B1 từ log B1 nếu có
        if LOG_B1.is_file():
            try:
                b1 = parse_metrics(LOG_B1)
                b1_vals = [b1["divider"], b1["ped_crossing"], b1["boundary"], b1["map"]]
                for i, (b2v, b1v) in enumerate(zip(b2_vals, b1_vals), start=1):
                    t.rows[i].cells[3].text = _fmt(b1v)
                    t.rows[i].cells[4].text = _fmt(b2v - b1v)
            except ValueError:
                pass

        for para in doc.paragraphs:
            if para.text.startswith("Ngày chạy B2:"):
                para.text = f"Ngày chạy B2 (lần gần nhất): {when}"

    # Lịch sử chạy (append)
    history = (
        f"[{when}] {step.upper()}: divider={_fmt(metrics['divider'])}, "
        f"ped={_fmt(metrics['ped_crossing'])}, boundary={_fmt(metrics['boundary'])}, "
        f"mAP={_fmt(metrics['map'])}"
    )
    doc.add_paragraph(history)

    doc.save(str(docx_path))


def ghi_ket_qua_b1(project_root: Path, metrics: dict[str, float], log_path: Path) -> None:
    when = datetime.now().strftime("%Y-%m-%d %H:%M")
    out = project_root / "outputs/nusc_eval/e1_camera"
    _write_metrics_md(out, "E1 — MapTRv2 camera", metrics, log_path)
    _update_docx(BAO_CAO_DOCX, step="b1", metrics=metrics, when=when)
    _update_docx(BAO_CAO_COPY, step="b1", metrics=metrics, when=when)


def ghi_ket_qua_b2(project_root: Path, metrics: dict[str, float], log_path: Path) -> None:
    when = datetime.now().strftime("%Y-%m-%d %H:%M")
    out = project_root / "outputs/nusc_eval/e2_fusion"
    _write_metrics_md(out, "E2 — MapTR fusion", metrics, log_path)
    _update_docx(BAO_CAO_DOCX, step="b2", metrics=metrics, when=when)
    _update_docx(BAO_CAO_COPY, step="b2", metrics=metrics, when=when)


def main() -> int:
    if len(sys.argv) < 2:
        print("Cách dùng: python scripts/ghi_bao_cao.py b1|b2")
        return 1

    step = sys.argv[1].lower()
    if step == "b1":
        metrics = parse_metrics(LOG_B1)
        ghi_ket_qua_b1(PROJECT_ROOT, metrics, LOG_B1)
    elif step == "b2":
        metrics = parse_metrics(LOG_B2)
        ghi_ket_qua_b2(PROJECT_ROOT, metrics, LOG_B2)
    else:
        print("Chỉ hỗ trợ: b1 hoặc b2")
        return 1

    print(f"Đã ghi vào {BAO_CAO_DOCX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
